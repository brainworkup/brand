#!/usr/bin/env python3
"""
generate_docx_refs.py
=====================
Reads each concept's _brand.yml and generates a branded reference.docx
that Quarto uses as a style template for docx output.

Usage:
    python3 generate_docx_refs.py              # all 7 concepts
    python3 generate_docx_refs.py concept-4-institute  # one concept

Produces: {concept}/reference.docx
"""

import sys
import os
import yaml
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

SCRIPT_DIR = Path(__file__).resolve().parent

CONCEPTS = [
    "concept-1-growing-minds",
    "concept-2-lighthouse",
    "concept-3-clinical-modern",
    "concept-4-institute",
    "concept-5-neural-blueprint",
    "concept-6-academic-authority",
    "concept-7-dual-identity",
]


def hex_to_rgb(hex_str: str) -> RGBColor:
    """Convert '#RRGGBB' to an RGBColor."""
    h = hex_str.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def resolve_color(value, palette: dict) -> str:
    """Resolve a color value: if it's a palette name, look it up; otherwise return as-is."""
    if isinstance(value, str) and not value.startswith("#"):
        return palette.get(value, value)
    return value


def load_brand(concept_dir: Path) -> dict:
    """Load and return the parsed _brand.yml with resolved colors."""
    brand_file = concept_dir / "_brand.yml"
    with open(brand_file) as f:
        brand = yaml.safe_load(f)

    # Build palette lookup
    palette = brand.get("color", {}).get("palette", {})

    # Resolve top-level color references
    color = brand.get("color", {})
    for key in ["foreground", "background", "primary", "secondary", "tertiary",
                "success", "info", "warning", "danger", "light", "dark"]:
        if key in color:
            color[key] = resolve_color(color[key], palette)

    # Resolve typography color references
    typo = brand.get("typography", {})
    for section in ["headings", "monospace-inline", "monospace-block", "link"]:
        if section in typo and "color" in typo[section]:
            typo[section]["color"] = resolve_color(typo[section]["color"], palette)
        if section in typo and "background-color" in typo[section]:
            typo[section]["background-color"] = resolve_color(
                typo[section]["background-color"], palette
            )

    return brand


def set_style_font(style, font_family: str, size_pt: float = None,
                   bold: bool = None, color: RGBColor = None):
    """Configure font properties on a Word style."""
    font = style.font
    font.name = font_family
    if size_pt is not None:
        font.size = Pt(size_pt)
    if bold is not None:
        font.bold = bold
    if color is not None:
        font.color.rgb = color

    # Set the east-asia and complex-script fonts too
    rpr = style.element.get_or_add_rPr()
    for tag in [qn("w:rFonts")]:
        el = rpr.find(tag)
        if el is None:
            el = parse_xml(f'<w:rFonts {nsdecls("w")} '
                           f'w:ascii="{font_family}" '
                           f'w:hAnsi="{font_family}" '
                           f'w:eastAsia="{font_family}" '
                           f'w:cs="{font_family}"/>')
            rpr.insert(0, el)
        else:
            el.set(qn("w:ascii"), font_family)
            el.set(qn("w:hAnsi"), font_family)
            el.set(qn("w:eastAsia"), font_family)
            el.set(qn("w:cs"), font_family)


def set_paragraph_spacing(style, before_pt: float = None, after_pt: float = None,
                          line_spacing: float = None):
    """Set paragraph spacing on a style."""
    pf = style.paragraph_format
    if before_pt is not None:
        pf.space_before = Pt(before_pt)
    if after_pt is not None:
        pf.space_after = Pt(after_pt)
    if line_spacing is not None:
        pf.line_spacing = line_spacing


def parse_size(size_str) -> float:
    """Convert a CSS-like size ('17px', '0.85em', '16px') to a pt value."""
    if isinstance(size_str, (int, float)):
        return float(size_str)
    s = str(size_str).strip().lower()
    if s.endswith("px"):
        return float(s[:-2]) * 0.75  # px to pt
    if s.endswith("pt"):
        return float(s[:-2])
    if s.endswith("em"):
        return float(s[:-2]) * 12  # relative to ~12pt base
    return 12.0


def generate_reference_docx(concept_name: str):
    """Generate a reference.docx for the given concept."""
    concept_dir = SCRIPT_DIR / concept_name
    brand = load_brand(concept_dir)

    color = brand.get("color", {})
    palette = color.get("palette", {})
    typo = brand.get("typography", {})

    # Extract key values
    base_font = typo.get("base", {}).get("family", "Calibri")
    base_size = parse_size(typo.get("base", {}).get("size", "11pt"))
    base_line_height = typo.get("base", {}).get("line-height", 1.15)

    heading_font = typo.get("headings", {}).get("family", base_font)
    heading_color_hex = resolve_color(
        typo.get("headings", {}).get("color", color.get("primary", "#000000")),
        palette,
    )
    heading_color = hex_to_rgb(heading_color_hex)

    mono_font = typo.get("monospace", {}).get("family", "Courier New")

    primary_hex = color.get("primary", "#0563C1")
    primary_color = hex_to_rgb(primary_hex)

    fg_hex = color.get("foreground", "#000000")
    fg_color = hex_to_rgb(fg_hex)

    link_color_hex = resolve_color(
        typo.get("link", {}).get("color", color.get("primary", "#0563C1")),
        palette,
    )
    link_color = hex_to_rgb(link_color_hex)

    # --- Build the document ---
    doc = Document()

    # -- Default style (Normal) --
    normal = doc.styles["Normal"]
    set_style_font(normal, base_font, base_size, color=fg_color)
    set_paragraph_spacing(normal, before_pt=0, after_pt=6, line_spacing=base_line_height)

    # -- Headings --
    heading_sizes = {
        "Heading 1": base_size * 2.0,
        "Heading 2": base_size * 1.6,
        "Heading 3": base_size * 1.3,
        "Heading 4": base_size * 1.1,
        "Heading 5": base_size * 1.0,
        "Heading 6": base_size * 0.9,
    }
    for style_name, size in heading_sizes.items():
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue
        set_style_font(style, heading_font, size, bold=True, color=heading_color)
        set_paragraph_spacing(style, before_pt=12, after_pt=4,
                              line_spacing=typo.get("headings", {}).get("line-height", 1.2))

    # -- Title --
    try:
        title_style = doc.styles["Title"]
        set_style_font(title_style, heading_font, base_size * 2.4, bold=True,
                       color=primary_color)
        set_paragraph_spacing(title_style, before_pt=0, after_pt=4, line_spacing=1.1)
    except KeyError:
        pass

    # -- Subtitle --
    try:
        subtitle_style = doc.styles["Subtitle"]
        set_style_font(subtitle_style, base_font, base_size * 1.3, bold=False,
                       color=heading_color)
        set_paragraph_spacing(subtitle_style, before_pt=0, after_pt=12, line_spacing=1.3)
    except KeyError:
        pass

    # -- Block Quote --
    try:
        bq = doc.styles["Quote"]
        set_style_font(bq, base_font, base_size, color=hex_to_rgb(
            color.get("tertiary", fg_hex)))
        bq.font.italic = True
        pf = bq.paragraph_format
        pf.left_indent = Inches(0.5)
    except KeyError:
        pass

    # -- Hyperlink character style --
    # Word doesn't always have a "Hyperlink" style accessible via python-docx,
    # but we can set it if it exists
    try:
        hl = doc.styles["Hyperlink"]
        set_style_font(hl, base_font, color=link_color)
    except KeyError:
        pass

    # -- Table styles: set the default table grid --
    try:
        tbl = doc.styles["Table Grid"]
        # We can't easily set table colors via python-docx styles,
        # but the font will carry through
    except KeyError:
        pass

    # -- Code / monospace styles --
    # Quarto uses "Source Code" style for code blocks in docx
    for code_style_name in ["Source Code", "Verbatim Char"]:
        try:
            cs = doc.styles[code_style_name]
            set_style_font(cs, mono_font, base_size * 0.85)
        except KeyError:
            pass

    # -- Add placeholder content so Word registers all styles --
    doc.add_heading("Title", level=0)  # Uses Title style
    doc.add_paragraph("Subtitle text", style="Subtitle")
    for level in range(1, 5):
        doc.add_heading(f"Heading {level}", level=level)
        doc.add_paragraph("Body text paragraph with normal styling.")
    doc.add_paragraph("A block quote.", style="Quote")

    # Code block placeholder
    try:
        doc.add_paragraph("code_example <- TRUE", style="Source Code")
    except KeyError:
        p = doc.add_paragraph("code_example <- TRUE")
        set_style_font(p.style, mono_font, base_size * 0.85)

    # Table placeholder
    table = doc.add_table(rows=2, cols=3, style="Table Grid")
    for i, cell in enumerate(table.rows[0].cells):
        cell.text = f"Header {i+1}"
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = primary_color
    for i, cell in enumerate(table.rows[1].cells):
        cell.text = f"Data {i+1}"

    # -- Save --
    out_path = concept_dir / "reference.docx"
    doc.save(str(out_path))
    print(f"  ✓ {concept_name}/reference.docx")


def main():
    filter_concept = None
    if len(sys.argv) > 1:
        filter_concept = sys.argv[1]

    concepts = [filter_concept] if filter_concept else CONCEPTS

    print(f"Generating {len(concepts)} reference.docx templates...\n")
    for concept in concepts:
        concept_dir = SCRIPT_DIR / concept
        if not (concept_dir / "_brand.yml").exists():
            print(f"  ✗ {concept}/_brand.yml not found, skipping")
            continue
        generate_reference_docx(concept)

    print(f"\nDone! Run: Rscript render_all_brands.R docx")


if __name__ == "__main__":
    main()
